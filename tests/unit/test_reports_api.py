import unittest
from app import create_app, db
from app.models import Employee
from datetime import datetime
import pandas as pd
from io import BytesIO
import openpyxl

class TestReportsAPI(unittest.TestCase):
    def setUp(self):
        self.app = create_app('testing')
        self.app_context = self.app.app_context()
        self.app_context.push()
        db.create_all()
        self.client = self.app.test_client()

        # Add dummy employees
        employees = [
            Employee(first_name='John', last_name='Doe', email='john.doe@example.com', department='Engineering', position='Developer', hire_date=datetime(2022, 1, 1), is_active=False),
            Employee(first_name='Jane', last_name='Smith', email='jane.smith@example.com', department='HR', position='Manager', hire_date=datetime(2021, 5, 15), is_active=True),
        ]
        db.session.add_all(employees)
        db.session.commit()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.app_context.pop()

    def test_export_report_csv(self):
        response = self.client.get('/analytics/api/reports/export?format=csv')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'text/csv')
        
        # Verify content
        data = BytesIO(response.data)
        df = pd.read_csv(data)
        self.assertEqual(len(df), 2)
        self.assertIn('john.doe@example.com', df['email'].values)

    def test_export_report_excel(self):
        response = self.client.get('/analytics/api/reports/export?format=excel')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        
        # Verify content
        data = BytesIO(response.data)
        workbook = openpyxl.load_workbook(data)
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 3) # Header + 2 rows
        self.assertEqual(sheet.cell(row=2, column=4).value, 'john.doe@example.com')

    def test_export_report_pdf(self):
        response = self.client.get('/analytics/api/reports/export?format=pdf')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'application/pdf')
        # PDF content verification is more complex, so we'll just check the mimetype and status code here.
        # A more thorough test could involve a library to parse the PDF content.

    def test_export_report_invalid_format(self):
        response = self.client.get('/analytics/api/reports/export?format=invalid')
        self.assertEqual(response.status_code, 400)
        self.assertIn('Invalid format specified', response.json['error'])

if __name__ == '__main__':
    unittest.main()
